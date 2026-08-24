import io
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and stamp 'Page X of Y' on all pages."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Times-Roman", 8)
        self.setFillColor(colors.HexColor("#333333"))
        
        # Formal Academic / Military Institutional Footer
        width, height = self._pagesize
        footer_left = "TTS Management Portal — Sri Lanka Air Force Trade Training School Ekala"
        page_str = f"Page {self._pageNumber} of {page_count}"
        security_text = "RESTRICTED / OFFICIAL USE ONLY"
        
        self.drawString(36, 20, footer_left)
        self.drawCentredString(width / 2.0, 20, security_text)
        self.drawRightString(width - 36, 20, page_str)
        
        # Thin single-line rule above footer
        self.setStrokeColor(colors.black)
        self.setLineWidth(0.5)
        self.line(36, 30, width - 36, 30)
        
        self.restoreState()


class ReportExportService:

    # ─────────────────────────────────────────────────────────────
    # 1. Native Excel Workbook Generator (.xlsx) — Classical Formal Style
    # ─────────────────────────────────────────────────────────────
    @staticmethod
    def generate_excel_workbook(
        report_title: str,
        header_info: Dict[str, Any],
        summary_data: Optional[Any],
        columns: List[Dict[str, Any]],
        rows: List[Dict[str, Any]]
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Official Report"
        ws.views.sheetView[0].showGridLines = True

        # Classical Institutional Palette
        dark_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        light_gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        title_font = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
        sub_font = Font(name="Times New Roman", size=11, bold=True, color="0F172A")
        meta_font = Font(name="Times New Roman", size=9.5, color="1E293B")
        meta_bold = Font(name="Times New Roman", size=9.5, bold=True, color="0F172A")
        tbl_hdr_font = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
        regular_font = Font(name="Times New Roman", size=9.5, color="000000")
        bold_font = Font(name="Times New Roman", size=9.5, bold=True, color="000000")
        
        thin_side = Side(border_style="thin", color="000000")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(bottom=Side(border_style="double", color="000000"))
        
        total_cols = max(len(columns), 6)

        # 1. Top Header Banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        c1 = ws.cell(row=1, column=1, value="SRI LANKA AIR FORCE — TRADE TRAINING SCHOOL (TTS EKALA)")
        c1.fill = dark_fill
        c1.font = title_font
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        c2 = ws.cell(row=2, column=1, value=report_title.upper())
        c2.font = sub_font
        c2.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # 2. Report Information / Parameters Table
        ws.cell(row=4, column=1, value="1. REPORT INFORMATION").font = sub_font
        ws.cell(row=4, column=1).border = double_bottom

        gen_at = header_info.get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        gen_by = header_info.get("generated_by", "Authorized System Officer")
        params = header_info.get("parameters", {})
        param_str = ", ".join([f"{k.replace('_', ' ').title()}: {v}" for k, v in params.items() if v]) or "All Master Records"

        info_rows = [
            ("Organization", "Sri Lanka Air Force (TTS Ekala)", "Generated Date/Time", str(gen_at)),
            ("Security Classification", "RESTRICTED / OFFICIAL USE ONLY", "Authorized Officer", str(gen_by)),
            ("Applied Parameters", str(param_str), "Total Records", f"{len(rows)} Records")
        ]

        cur_r = 5
        for r_data in info_rows:
            ws.cell(row=cur_r, column=1, value=r_data[0]).font = meta_bold
            ws.cell(row=cur_r, column=2, value=r_data[1]).font = meta_font
            ws.cell(row=cur_r, column=4, value=r_data[2]).font = meta_bold
            ws.cell(row=cur_r, column=5, value=r_data[3]).font = meta_font
            cur_r += 1

        cur_r += 1

        # 3. Statistical Summary Table (if available)
        if summary_data:
            ws.cell(row=cur_r, column=1, value="2. EXECUTIVE STATISTICAL SUMMARY").font = sub_font
            ws.cell(row=cur_r, column=1).border = double_bottom
            cur_r += 1

            sum_items = []
            if isinstance(summary_data, dict):
                sum_items = list(summary_data.items())
            elif isinstance(summary_data, list):
                sum_items = [(item.get("label", ""), item.get("value", "")) for item in summary_data if isinstance(item, dict)]

            if sum_items:
                ws.cell(row=cur_r, column=1, value="Parameter / Metric").font = tbl_hdr_font
                ws.cell(row=cur_r, column=1).fill = dark_fill
                ws.cell(row=cur_r, column=2, value="Value / Rate").font = tbl_hdr_font
                ws.cell(row=cur_r, column=2).fill = dark_fill
                cur_r += 1

                for k, v in sum_items:
                    c_k = ws.cell(row=cur_r, column=1, value=str(k).replace("_", " ").title())
                    c_k.font = bold_font
                    c_k.border = thin_border
                    c_k.fill = light_gray_fill

                    c_v = ws.cell(row=cur_r, column=2, value=str(v))
                    c_v.font = regular_font
                    c_v.border = thin_border
                    c_v.alignment = Alignment(horizontal="right")
                    cur_r += 1

                cur_r += 1

        # 4. Detailed Records Table
        sec_num = "3." if summary_data else "2."
        ws.cell(row=cur_r, column=1, value=f"{sec_num} DETAILED SYSTEM RECORDS").font = sub_font
        ws.cell(row=cur_r, column=1).border = double_bottom
        cur_r += 1

        # Table Headers
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=cur_r, column=col_idx, value=col.get("label", col.get("field")).upper())
            cell.fill = dark_fill
            cell.font = tbl_hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[cur_r].height = 22
        cur_r += 1

        # Table Rows
        for r_idx, r in enumerate(rows):
            is_even = (r_idx % 2 == 0)
            for col_idx, col in enumerate(columns, start=1):
                field = col.get("field")
                val = r.get(field, "—")
                cell = ws.cell(row=cur_r, column=col_idx, value=str(val) if val is not None else "—")
                cell.font = regular_font
                cell.border = thin_border
                if not is_even:
                    cell.fill = zebra_fill
                
                align = col.get("align", "left")
                cell.alignment = Alignment(horizontal=align, vertical="center")
            cur_r += 1

        cur_r += 2

        # 5. Signatures Block
        ws.cell(row=cur_r, column=1, value="PREPARED BY:").font = bold_font
        ws.cell(row=cur_r, column=3, value="CHECKED & VERIFIED BY:").font = bold_font
        ws.cell(row=cur_r, column=5, value="AUTHENTICATED BY:").font = bold_font
        cur_r += 3
        ws.cell(row=cur_r, column=1, value="..................................................").font = meta_font
        ws.cell(row=cur_r, column=3, value="..................................................").font = meta_font
        ws.cell(row=cur_r, column=5, value="..................................................").font = meta_font
        cur_r += 1
        ws.cell(row=cur_r, column=1, value="Signature / Date").font = meta_font
        ws.cell(row=cur_r, column=3, value="Signature / Date").font = meta_font
        ws.cell(row=cur_r, column=5, value="Signature / Date").font = meta_font

        # Auto-fit Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 36), 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    # ─────────────────────────────────────────────────────────────
    # 2. Native PDF Document Generator (.pdf) — Classical Academic / Institutional
    # ─────────────────────────────────────────────────────────────
    @staticmethod
    def generate_pdf_document(
        report_title: str,
        header_info: Dict[str, Any],
        summary_data: Optional[Any],
        columns: List[Dict[str, Any]],
        rows: List[Dict[str, Any]]
    ) -> bytes:
        output = io.BytesIO()

        # Orientation selection: landscape for wide reports (> 6 columns or academic/attendance/calendar/parade)
        is_wide = len(columns) > 6 or any(k in report_title.lower() for k in ["academic", "attendance", "calendar", "parade", "marksheet"])
        page_size = landscape(A4) if is_wide else portrait(A4)

        doc = SimpleDocTemplate(
            output,
            pagesize=page_size,
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=38
        )

        styles = getSampleStyleSheet()
        
        # Classical Typography Styles (Times-Roman / Times-Bold)
        org_title_style = ParagraphStyle(
            'OrgTitle',
            fontName='Times-Bold',
            fontSize=12,
            leading=14,
            textColor=colors.black,
            alignment=1  # Centered
        )
        school_style = ParagraphStyle(
            'SchoolTitle',
            fontName='Times-Bold',
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor('#1E293B'),
            alignment=1
        )
        report_title_style = ParagraphStyle(
            'ReportTitleStyle',
            fontName='Times-Bold',
            fontSize=11.5,
            leading=15,
            textColor=colors.black,
            alignment=1
        )
        section_heading_style = ParagraphStyle(
            'SectionHeading',
            fontName='Times-Bold',
            fontSize=9.5,
            leading=12,
            textColor=colors.black
        )
        meta_label_style = ParagraphStyle(
            'MetaLabel',
            fontName='Times-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.black
        )
        meta_val_style = ParagraphStyle(
            'MetaValue',
            fontName='Times-Roman',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#1E293B')
        )
        tbl_hdr_style = ParagraphStyle(
            'TableHdr',
            fontName='Times-Bold',
            fontSize=7.5,
            leading=9,
            textColor=colors.black,
            alignment=1
        )
        cell_regular_style = ParagraphStyle(
            'CellRegular',
            fontName='Times-Roman',
            fontSize=7.5,
            leading=9,
            textColor=colors.black
        )
        cell_center_style = ParagraphStyle(
            'CellCenter',
            fontName='Times-Roman',
            fontSize=7.5,
            leading=9,
            textColor=colors.black,
            alignment=1
        )
        cell_bold_style = ParagraphStyle(
            'CellBold',
            fontName='Times-Bold',
            fontSize=7.5,
            leading=9,
            textColor=colors.black
        )

        elements = []

        # 1. Classical Institutional Header
        elements.append(Paragraph("SRI LANKA AIR FORCE", org_title_style))
        elements.append(Paragraph("TRADE TRAINING SCHOOL — EKALA", school_style))
        elements.append(Spacer(1, 4))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=4))
        elements.append(Paragraph(report_title.upper(), report_title_style))
        elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.black, spaceBefore=4, spaceAfter=8))

        # 2. Report Information Section (Structured 2-column Parameter Block)
        gen_at = header_info.get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        gen_by = header_info.get("generated_by", "Authorized System Officer")
        params = header_info.get("parameters", {})
        param_str = ", ".join([f"{k.replace('_', ' ').title()}: {v}" for k, v in params.items() if v]) or "All Master Records"

        info_data = [
            [
                Paragraph("<b>Organization:</b>", meta_label_style),
                Paragraph("Sri Lanka Air Force (TTS Ekala)", meta_val_style),
                Paragraph("<b>Generated Date/Time:</b>", meta_label_style),
                Paragraph(str(gen_at), meta_val_style)
            ],
            [
                Paragraph("<b>Security Classification:</b>", meta_label_style),
                Paragraph("RESTRICTED / OFFICIAL USE ONLY", meta_val_style),
                Paragraph("<b>Authorized Officer:</b>", meta_label_style),
                Paragraph(str(gen_by), meta_val_style)
            ],
            [
                Paragraph("<b>Parameters:</b>", meta_label_style),
                Paragraph(str(param_str), meta_val_style),
                Paragraph("<b>Total Records:</b>", meta_label_style),
                Paragraph(f"<b>{len(rows)} Records</b>", meta_val_style)
            ]
        ]
        
        info_table = Table(info_data, colWidths=['16%', '34%', '18%', '32%'])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('PADDING', (0,0), (-1,-1), 3),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.append(Paragraph("1. REPORT INFORMATION", section_heading_style))
        elements.append(Spacer(1, 3))
        elements.append(info_table)
        elements.append(Spacer(1, 8))

        # 3. Summary Statistics Section (Table 1: Statistical Summary)
        if summary_data:
            sum_items = []
            if isinstance(summary_data, dict):
                sum_items = list(summary_data.items())
            elif isinstance(summary_data, list):
                sum_items = [(item.get("label", ""), item.get("value", "")) for item in summary_data if isinstance(item, dict)]

            if sum_items:
                elements.append(Paragraph("2. STATISTICAL SUMMARY", section_heading_style))
                elements.append(Spacer(1, 3))

                sum_rows = [
                    [Paragraph("<b>Metric / Parameter</b>", tbl_hdr_style), Paragraph("<b>Count / Value</b>", tbl_hdr_style),
                     Paragraph("<b>Metric / Parameter</b>", tbl_hdr_style), Paragraph("<b>Count / Value</b>", tbl_hdr_style)]
                ]

                # Pair metrics into 2-column pairs
                for i in range(0, len(sum_items), 2):
                    p1_lbl, p1_val = sum_items[i]
                    p2_lbl, p2_val = sum_items[i+1] if (i+1) < len(sum_items) else ("", "")
                    sum_rows.append([
                        Paragraph(str(p1_lbl).replace("_", " ").title(), cell_bold_style),
                        Paragraph(str(p1_val), cell_center_style),
                        Paragraph(str(p2_lbl).replace("_", " ").title(), cell_bold_style) if p2_lbl else Paragraph("", cell_regular_style),
                        Paragraph(str(p2_val), cell_center_style) if p2_lbl else Paragraph("", cell_regular_style)
                    ])

                sum_table = Table(sum_rows, colWidths=['32%', '18%', '32%', '18%'])
                sum_table.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.black),
                    ('PADDING', (0,0), (-1,-1), 2.5),
                    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ]))
                elements.append(sum_table)
                elements.append(Spacer(1, 8))

        # 4. Detailed Records Section (Table 2: Detailed System Register)
        sec_title = "3. DETAILED SYSTEM RECORDS" if summary_data else "2. DETAILED SYSTEM RECORDS"
        elements.append(Paragraph(sec_title, section_heading_style))
        elements.append(Spacer(1, 3))

        table_data = []
        hdr_row = [Paragraph(f"<b>{col.get('label', col.get('field')).upper()}</b>", tbl_hdr_style) for col in columns]
        table_data.append(hdr_row)

        for r in rows:
            row_cells = []
            for col in columns:
                field = col.get("field")
                val = str(r.get(field, "—")) if r.get(field) is not None else "—"
                align = col.get("align", "left")
                
                if align == "center":
                    st = cell_center_style
                elif field in ["service_number", "s_no", "status", "result_status"]:
                    st = cell_bold_style
                else:
                    st = cell_regular_style

                row_cells.append(Paragraph(val, st))
            table_data.append(row_cells)

        usable_w = doc.pagesize[0] - 48
        col_w = usable_w / len(columns)

        data_table = Table(table_data, colWidths=[col_w] * len(columns), repeatRows=1)
        data_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.black),
            ('PADDING', (0,0), (-1,-1), 2.5),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.append(data_table)

        # 5. Formal Certification & Signature Section
        sig_data = [
            [
                Paragraph("<b>PREPARED BY:</b>", meta_label_style),
                Paragraph("<b>CHECKED & VERIFIED BY:</b>", meta_label_style),
                Paragraph("<b>APPROVED & AUTHENTICATED BY:</b>", meta_label_style)
            ],
            [
                Paragraph("<br/><br/>........................................................", meta_val_style),
                Paragraph("<br/><br/>........................................................", meta_val_style),
                Paragraph("<br/><br/>........................................................", meta_val_style)
            ],
            [
                Paragraph("Name: ............................................<br/>Service No: ....................................<br/>Designation: System Officer<br/>Date: .............................................", meta_val_style),
                Paragraph("Name: ............................................<br/>Service No: ....................................<br/>Designation: Chief Ground Instructor<br/>Date: .............................................", meta_val_style),
                Paragraph("Name: ............................................<br/>Service No: ....................................<br/>Designation: Commanding Officer / OC<br/>Date: .............................................", meta_val_style)
            ]
        ]
        sig_table = Table(sig_data, colWidths=['33.3%', '33.3%', '33.4%'])
        sig_table.setStyle(TableStyle([
            ('PADDING', (0,0), (-1,-1), 2),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))

        elements.append(Spacer(1, 14))
        elements.append(KeepTogether([
            Paragraph("OFFICIAL AUTHENTICATION & CERTIFICATION", section_heading_style),
            HRFlowable(width="100%", thickness=0.5, color=colors.black, spaceBefore=2, spaceAfter=6),
            sig_table
        ]))

        doc.build(elements, canvasmaker=NumberedCanvas)
        output.seek(0)
        return output.getvalue()


report_export_service = ReportExportService()
